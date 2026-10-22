from __future__ import annotations
# -*- coding: utf-8 -*-
"""Excel-отчёт: SUMMARY + DATA_QUALITY + по листу на проверку."""
import os
from datetime import datetime

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.engine import RunResult

FILLS = {"OK": PatternFill("solid", fgColor="C6EFCE"),
         "ISSUE": PatternFill("solid", fgColor="FFC7CE"),
         "REVIEW": PatternFill("solid", fgColor="FFEB9C"),
         "SKIPPED": PatternFill("solid", fgColor="D9D9D9"),
         "ERROR": PatternFill("solid", fgColor="FFC7CE"),
         "WARNING": PatternFill("solid", fgColor="FFEB9C"),
         "INFO": PatternFill("solid", fgColor="DDEBF7")}
FONTS = {"OK": Font(color="006100"), "ISSUE": Font(color="9C0006", bold=True),
         "REVIEW": Font(color="9C6500"), "SKIPPED": Font(color="595959"),
         "ERROR": Font(color="9C0006", bold=True), "WARNING": Font(color="9C6500"),
         "INFO": Font(color="1F4E79")}


def write_excel(result: RunResult, cfg: dict, out_path: str) -> str:
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    summary = pd.DataFrame(result.summary_rows())
    if result.findings:
        dq = pd.DataFrame([f.as_dict() for f in result.findings])
    else:
        dq = pd.DataFrame(columns=["severity", "source", "section", "column", "issue",
                                   "affected_rows", "examples", "impact"])
    src = pd.DataFrame(result.sources_info)
    meta = pd.DataFrame([
        {"key": "generated_at", "value": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
        {"key": "duration_sec", "value": result.duration_sec},
        {"key": "total_issues", "value": result.total_issues()},
        {"key": "dq_errors", "value": result.dq_counts()["ERROR"]},
        {"key": "dq_warnings", "value": result.dq_counts()["WARNING"]},
    ])

    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        meta.to_excel(xw, sheet_name="SUMMARY", index=False, startrow=0)
        summary.to_excel(xw, sheet_name="SUMMARY", index=False, startrow=len(meta) + 2)
        src.to_excel(xw, sheet_name="SOURCES", index=False)
        dq.to_excel(xw, sheet_name="DATA_QUALITY", index=False)
        for cid, out in result.outputs.items():
            sheet = f"{cid}"[:31]
            out.df.to_excel(xw, sheet_name=sheet, index=False)

        wb = xw.book
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            for c, col in enumerate(ws.columns, 1):
                width = min(70, max(12, max((len(str(cell.value or "")) for cell in col), default=12) + 2))
                ws.column_dimensions[get_column_letter(c)].width = width
                for cell in col:
                    cell.alignment = Alignment(vertical="center")
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1A56DB")
            hdr = [cell.value for cell in ws[1]]
            for key in ("status", "severity"):
                if key in hdr:
                    si = hdr.index(key) + 1
                    for row in ws.iter_rows(min_row=2, min_col=si, max_col=si):
                        cell = row[0]
                        if cell.value in FILLS:
                            cell.fill = FILLS[cell.value]
                            cell.font = FONTS[cell.value]
    return out_path
