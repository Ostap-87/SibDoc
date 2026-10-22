from __future__ import annotations
# -*- coding: utf-8 -*-
"""Вставка извлечённых из сканов данных в Excel для предпросмотра.

Маппинг поле -> колонка листа SCAN_PREVIEW живёт в
config/scan_mapping.yaml и меняется без кода.
"""
import os
from typing import Any

import pandas as pd
import yaml
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))

STATUS_FILL = {
    "found":   PatternFill("solid", fgColor="C6EFCE"),
    "review":  PatternFill("solid", fgColor="FFEB9C"),
    "missing": PatternFill("solid", fgColor="FFC7CE"),
}


def load_mapping() -> dict:
    path = os.path.join(PROJECT_ROOT, "config", "scan_mapping.yaml")
    with open(path, encoding="utf-8") as f:
        return (yaml.safe_load(f) or {}).get("target", {})


def build_preview_rows(scan_results: list[dict]) -> list[dict]:
    """scan_results — список {file, fields: {field: {value,status,...}}}.
    Возвращает строки для UI-таблицы предпросмотра."""
    rows = []
    for res in scan_results:
        fields = res.get("fields", {})
        statuses = [f.get("status") for f in fields.values()]
        if "review" in statuses:
            overall = "review"
        elif all(s == "missing" for s in statuses):
            overall = "missing"
        elif "missing" in statuses:
            overall = "review"
        else:
            overall = "found"
        row = {"file": res.get("file", ""), "status": overall,
               "pages": res.get("pages", 1), "langs": "+".join(res.get("langs", []))}
        for field, rec in fields.items():
            row[field] = rec.get("value", "")
            row[f"_{field}_status"] = rec.get("status", "missing")
        rows.append(row)
    return rows


def write_preview_excel(scan_results: list[dict], out_path: str) -> str:
    """Сформировать Excel предпросмотра: извлечённые значения в нужных
    ячейках, раскраска по статусу извлечения."""
    mapping = load_mapping()
    cols: dict = mapping.get("columns", {})
    svc: dict = mapping.get("service_columns", {})

    order = list(cols.keys())
    header = [cols[f] for f in order] + [svc.get("file", "Source file"),
                                         svc.get("status", "Extract status")]

    data_rows = []
    for res in scan_results:
        fields = res.get("fields", {})
        statuses = [f.get("status") for f in fields.values()]
        overall = ("missing" if all(s == "missing" for s in statuses)
                   else "review" if ("review" in statuses or "missing" in statuses)
                   else "found")
        data_rows.append([fields.get(f, {}).get("value", "") for f in order]
                         + [res.get("file", ""), overall])
        data_rows[-1] = data_rows[-1]  # keep flat

    df = pd.DataFrame(data_rows, columns=header)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name=mapping.get("sheet", "SCAN_PREVIEW"), index=False)
        ws = xw.book[mapping.get("sheet", "SCAN_PREVIEW")]
        ws.freeze_panes = "A2"
        for c, col in enumerate(ws.columns, 1):
            width = min(40, max(14, max((len(str(cell.value or "")) for cell in col), default=14) + 2))
            ws.column_dimensions[get_column_letter(c)].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1A56DB")
        # раскраска ячеек по статусу извлечения
        for ri, res in enumerate(scan_results, start=2):
            fields = res.get("fields", {})
            for ci, field in enumerate(order, start=1):
                st = fields.get(field, {}).get("status", "missing")
                ws.cell(row=ri, column=ci).fill = STATUS_FILL.get(st, STATUS_FILL["missing"])
            overall_cell = ws.cell(row=ri, column=len(header))
            ov = overall_cell.value
            if ov in STATUS_FILL:
                overall_cell.fill = STATUS_FILL[ov]
                overall_cell.font = Font(bold=True)
    return out_path
