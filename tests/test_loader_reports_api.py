# -*- coding: utf-8 -*-
"""Загрузка файлов (кодировки, разделители, листы, XML, ZIP, словарь полей), отчёты, API."""
from __future__ import annotations

import glob
import io
import os
import zipfile

import openpyxl
import pandas as pd
import pytest

from app.services.loader import LoadError, read_many, read_table
from app.services.xml_decl import XmlDeclError, parse_xml

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))


# ------------------------------------------------------------------ загрузка
def test_csv_gbk_and_semicolon_and_cp1251(cfg):
    df = pd.DataFrame({"集装箱号": ["DEMU0001012"], "海关编号": ["220220260000000101"], "单价": ["10,157"]})
    gbk = df.to_csv(index=False).encode("gb18030")
    out, info = read_table(gbk, "cn.csv", "cn_dec", cfg)
    assert info["encoding"] == "gb18030" and {"container", "decl_cn", "unit_price"} <= set(out.columns)
    ru = "Номер ДТ;Контейнер;Вес нетто\n10000010/150926/0000001;DEMU0001012;23 940,5\n".encode("cp1251")
    out, info = read_table(ru, "dt.csv", "ru_dt", cfg)
    assert info["delimiter"] == ";" and info["encoding"] == "cp1251"
    assert out.loc[0, "decl_ru"] == "10000010/150926/0000001" and "net_weight" in out.columns


def test_excel_sheet_and_header_row_detection(cfg, tmp_path):
    path = tmp_path / "p.xlsx"
    with pd.ExcelWriter(path) as xw:
        pd.DataFrame({"x": [1]}).to_excel(xw, sheet_name="Обложка", index=False)
        pd.DataFrame([["Выгрузка SAP за октябрь", None, None, None, None],
                      ["Container number", "Price", "Main contract", "Add contract", "Net weight"],
                      ["DEMU0001012", 10157, "ДЕМО.1001", "01А", 23.94]]).to_excel(
            xw, sheet_name="Данные", index=False, header=False)
    df, info = read_table(path.read_bytes(), "p.xlsx", "purchase", cfg)
    assert info["sheet"] == "Данные" and info["header_row"] == 2 and df.loc[0, "Price"] == 10157


def test_column_synonyms_to_expected_names(cfg):
    df = pd.DataFrame({"Номер контейнера": ["DEMU0001012"], "Цена SAP": [10157], "Основной контракт": ["ДЕМО.1001"],
                       "Доп. соглашение": ["01А"], "Вес нетто": [23.94]})
    out, info = read_table(df.to_csv(index=False).encode("utf-8"), "p.csv", "purchase", cfg)
    assert {"Container number", "Price", "Main contract", "Add contract", "Net weight"} <= set(out.columns)
    assert len(info["mapped"]) == 5


def test_ru_dt_xml_cp1251(cfg, demo_dir):
    f = sorted(glob.glob(os.path.join(demo_dir, "ru_dt", "*.xml")))[0]
    df, meta = parse_xml(open(f, "rb").read(), os.path.basename(f))
    r = df.iloc[0]
    assert meta["slot"] == "ru_dt" and r["decl_ru"] == "10000010/150926/0000001"
    assert r["contract"] == "ДЕМО.1001-01А" and r["container"].count("DEMU") == 2
    assert r["net_weight"] == 47880 and r["currency"] == "CNY" and r["incoterms"] == "CPT"


def test_cn_dec_xml_codes(cfg, demo_dir):
    f = sorted(glob.glob(os.path.join(demo_dir, "cn_dec_xml", "*.xml")))[2]
    df, meta = parse_xml(open(f, "rb").read(), os.path.basename(f))
    r = df.iloc[0]
    assert meta["slot"] == "cn_dec" and r["currency"] == "CNY"          # 142 -> CNY
    assert r["incoterms"] == "CFR" and r["incoterms_code"] == "2"      # 成交方式 2 = C&F
    assert r["unit"] == "kg" and r["decl_date"] == "2026-09-24" and r["net_weight"] > 0


def test_xml_rejects_entities():
    bad = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]><ESADout_CU>&a;</ESADout_CU>'
    with pytest.raises(XmlDeclError):
        parse_xml(bad, "x.xml")


def test_zip_of_xml(cfg, demo_dir):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for f in glob.glob(os.path.join(demo_dir, "ru_dt", "*.xml")):
            z.write(f, "dt/" + os.path.basename(f))
    df, info = read_table(buf.getvalue(), "all.zip", "ru_dt", cfg)
    assert info["declarations"] == 6 and df["decl_ru"].nunique() == 6


def test_unknown_format(cfg):
    with pytest.raises(LoadError):
        read_table(b"abc", "x.docx", "wide_table", cfg)


# ------------------------------------------------------------------ отчёты
def test_reports_formula_escape_and_sheets(demo_sources, cfg, tmp_path, offline_rates):
    from app.services.engine import run
    from app.services.report_excel import write_excel
    from app.services.report_pdf import write_pdf
    demo_sources["wide_table"].df.loc[0, "product_name"] = "=HYPERLINK(\"http://x\",\"y\")"
    res = run(demo_sources, cfg)
    x = write_excel(res, cfg, str(tmp_path / "r.xlsx"))
    wb = openpyxl.load_workbook(x)
    assert {"Сводка", "ДТ РФ", "报关单 КНР", "6 ДТ-报关单-SAP"} <= set(wb.sheetnames)
    formulas = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.data_type == "f"]
    assert not formulas
    p = write_pdf(res, cfg, str(tmp_path / "r.pdf"))
    import pymupdf
    text = "".join(pg.get_text() for pg in pymupdf.open(p))
    assert "报关单" in text and "ДТ" in text and "Сводка" in text


# ------------------------------------------------------------------ API
@pytest.fixture()
def client(isolated_storage, monkeypatch, tmp_path):
    from fastapi.testclient import TestClient
    import app.server as server
    monkeypatch.setattr(server, "UPLOADS_DIR", str(tmp_path / "uploads"))
    server._uploads.clear()
    return TestClient(server.app), server


def test_api_upload_run_and_journal(client, demo_dir):
    c, server = client
    files = [("files", (os.path.basename(f), open(f, "rb").read()))
             for f in sorted(glob.glob(os.path.join(demo_dir, "ru_dt", "*.xml")))]
    assert c.post("/api/upload/ru_dt", files=files).json()["rows"] == 6
    for slot, f in (("wide_table", "wide_table_demo.xlsx"), ("purchase", "purchase_demo.xlsx")):
        assert c.post(f"/api/upload/{slot}", files={"file": (f, open(os.path.join(demo_dir, f), "rb").read())}).status_code == 200
    d = c.get("/api/declarations/ru_dt").json()
    assert d["count"] == 6 and d["declarations"][0]["goods"]
    r = c.post("/api/run").json()
    assert r["run_id"] and c.get(r["reports"]["excel"]).status_code == 200 and c.get(r["reports"]["pdf"]).status_code == 200
    kinds = [e["kind"] for e in c.get("/api/journal").json()["entries"]]
    assert {"upload", "run", "report"} <= set(kinds)
    assert c.get("/api/journal/export").content.startswith("﻿".encode())


def test_api_replace_does_not_resurrect_old_file(client, cfg, tmp_path):
    c, server = client
    a = pd.DataFrame({"decl_num": ["1", "2", "3"]}).to_csv(index=False).encode()
    b = pd.DataFrame({"decl_num": ["7"]})
    xb = io.BytesIO()
    b.to_excel(xb, index=False)
    c.post("/api/upload/sap_vs_cd", files={"file": ("m.csv", a)})
    c.post("/api/upload/sap_vs_cd", files={"file": ("m.xlsx", xb.getvalue())})
    server._uploads.clear()
    server._restore_uploads()
    assert server._uploads["sap_vs_cd"].rows == 1


def test_api_rules_journal_and_history(client):
    c, server = client
    r = c.get("/api/rules/price.yaml").json()["rules"]
    r["tolerance_abs"] = 5
    d = c.put("/api/rules/price.yaml", json=r).json()
    assert d["diff"] == [{"key": "tolerance_abs", "old": 0.01, "new": 5}]
    e = c.get("/api/journal?kind=rules_change").json()["entries"][0]
    assert "tolerance_abs" in e["detail"]
    hist = c.get("/api/rules/price.yaml").json()["history"]
    assert hist and c.post(f"/api/rules/price.yaml/restore/{hist[0]['version']}").status_code == 200
    assert c.get("/api/rules/evil.yaml").status_code == 400
    assert c.put("/api/rules/evil.yaml", json={}).status_code == 400


def test_api_operator_cannot_change_rules(client, monkeypatch):
    c, server = client
    monkeypatch.setattr(server, "_auth_disabled", lambda: False)
    from app import security
    monkeypatch.setattr(security, "get_session", lambda t: {"login": "op", "role": "operator"} if t else None)
    r = c.put("/api/rules/price.yaml", json={"tolerance_abs": 1}, headers={"Authorization": "Bearer x"})
    assert r.status_code == 403


def test_api_upload_size_limit(client, monkeypatch):
    c, server = client
    monkeypatch.setattr(server, "_max_bytes", lambda cfg: 10)
    assert c.post("/api/upload/sap_vs_cd", files={"file": ("m.csv", b"x" * 100)}).status_code == 413
